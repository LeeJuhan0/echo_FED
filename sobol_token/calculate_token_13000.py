import math
import os
import re
import glob
import argparse
from typing import List, Dict, Tuple, Optional

import numpy as np
import pandas as pd
from neo4j import GraphDatabase
import tiktoken
from openpyxl import load_workbook

# -------------------------------
# Configurable parameters
# -------------------------------

EXCEL_PATH = r"C:\Users\HUFS_MATH\IdeaProjects\FOMC_Graphrag\Simscore관련\openai_paragraph_vs_statements_only_economic.xlsx"
EXCEL_ROW_INDEX = 1
EXCEL_START_COL_INDEX = 1
EXCEL_GID_COUNT = 70

OPENPYXL_FIRST_GID_COL = 2
OPENPYXL_MIN_ROW = 3
OPENPYXL_MAX_ROW = 2044
OPENPYXL_SHEETS_SLICE = slice(0, 2)

NEO4J_URI = os.getenv("NEO4J_URI", "bolt://localhost:7687")
NEO4J_USER = os.getenv("NEO4J_USER", "neo4j")
NEO4J_PASSWORD = os.getenv("NEO4J_PASSWORD", "12345678")

SIM_SCORE_INDEX = 0

MODEL_NAME = "gpt-5-2025-08-07"
TOKEN_TOLERANCE = 128

SYS_PROMPT_TWO = """
Modify the response to the question by incorporating insights from provided financial books or academic papers. 
Should adjust the sentiment score accordingly, ensuring that the assessment reflects provided literature in economics and finance.

The output must end with the sentiment score in the following exact format: "sentiment score = %float"
""".strip()

LAST_RESPONSE = '-0.40 "The effects of the coronavirus will weigh on economic activity in the near term and pose risks to the economic outlook." sentiment score = -0.40'

# 고정 토큰수(모든 샘플에 대해 동일)
FIXED_TARGET_TOKENS = 19000
FIXED_BASE_TOKENS = 1300

PERCENT_START = 0.0
PERCENT_END = 100.0
PERCENT_STEP = 0.1

REPORT_SUM_DIR = r"C:\Users\HUFS_MATH\IdeaProjects\FOMC_Graphrag\Statement\report_sum"

# Debug controls
DEBUG = True
DEBUG_MAX_P_LOGS_PER_SAMPLE = 200
DEBUG_PRINT_CONT_LEN_ONLY = True

# New options
USE_GRAPH_PERCENTILE = True   # True: percentile from Neo4j scores; False: from Excel column
COMPARISON_OP = ">="          # ">=" recommended to avoid strict boundary drops (use ">" if needed)
THRESHOLD_EPS = 1e-12         # subtract small epsilon from threshold to include boundary ties when using ">"

# -------------------------------
# Tokenization helpers
# -------------------------------

def encoding_for_model_safe(model: str):
    try:
        return tiktoken.encoding_for_model(model)
    except Exception:
        raise Exception("모델 정보 불일치")
        if "gpt-4o" in model or "o200k" in model:
            return tiktoken.get_encoding("o200k_base")
        return tiktoken.get_encoding("cl100k_base")

def count_chat_tokens(messages: List[Dict[str, str]], model: str) -> int:
    enc = encoding_for_model_safe(model)
    tokens_per_message = 3
    tokens_per_name = 1
    num_tokens = 0
    for msg in messages:
        num_tokens += tokens_per_message
        for k, v in msg.items():
            if k == "content":
                num_tokens += len(enc.encode(v))
            elif k == "name":
                num_tokens += tokens_per_name
    num_tokens += 3
    return num_tokens

# -------------------------------
# GID from Excel
# -------------------------------

def load_gids_from_excel(path: str, row_index: int, start_col_index: int, count: int) -> Tuple[List[str], Dict[str, int]]:
    df = pd.read_excel(path, header=None)
    series = df.iloc[row_index, start_col_index:start_col_index + count]
    series = series.dropna()
    gids: List[str] = []
    for v in series.tolist():
        s = str(int(v)) if isinstance(v, (int, float)) and not math.isnan(v) else str(v)
        s = s.strip().replace("-", "").replace(".", "")
        s = "".join(ch for ch in s if ch.isdigit())
        if s:
            gids.append(f"FOMC{s}")
    if len(gids) != count:
        print(f"[WARN] Loaded {len(gids)} gids, expected {count}. Proceeding with loaded values.")
    gid_to_excel_col: Dict[str, int] = {gid: OPENPYXL_FIRST_GID_COL + idx for idx, gid in enumerate(gids)}
    return gids, gid_to_excel_col

# -------------------------------
# report_sum -> question
# -------------------------------

def extract_yyyymm_from_gid(gid: str) -> Optional[str]:
    m = re.search(r"FOMC(\d{6})", gid)
    return m.group(1) if m else None

def pick_best_csv_match(candidates: List[str]) -> Optional[str]:
    if not candidates:
        return None
    dated = []
    for p in candidates:
        m = re.search(r"(\d{8})", os.path.basename(p))
        if m:
            try:
                dated.append((int(m.group(1)), p))
            except ValueError:
                pass
    if dated:
        dated.sort(key=lambda x: x[0], reverse=True)
        return dated[0][1]
    candidates.sort()
    return candidates[0]

def find_report_csv_for_gid(report_dir: str, gid: str) -> Optional[str]:
    yyyymm = extract_yyyymm_from_gid(gid)
    if not yyyymm:
        return None
    patterns = [
        os.path.join(report_dir, f"monetary{yyyymm}*.csv"),
        os.path.join(report_dir, f"*{yyyymm}*.csv"),
        os.path.join(report_dir, f"*{yyyymm}*monetary*.csv"),
    ]
    candidates = []
    for pat in patterns:
        candidates.extend(glob.glob(pat))
    candidates = list(set(candidates))
    return pick_best_csv_match(candidates)

def read_file_text(path: str, encoding: str = "utf-8") -> Optional[str]:
    try:
        with open(path, "r", encoding=encoding, errors="ignore") as f:
            return f.read()
    except Exception as e:
        print(f"[WARN] Failed to read file: {path} ({e})")
        return None

def get_question_for_gid(gid: str) -> str:
    csv_path = find_report_csv_for_gid(REPORT_SUM_DIR, gid)
    if csv_path:
        text = read_file_text(csv_path)
        if text and text.strip():
            return text.strip()
    # 명시적 폴백 제거 (질문 텍스트가 꼭 필요 없다면 빈 문자열로도 가능)
    return ""

# -------------------------------
# Excel percentile helpers
# -------------------------------

def preload_excel_columns(file_path: str, col_indices: List[int], min_row: int, max_row: int, sheets_slice: slice) -> Dict[int, np.ndarray]:
    wb = load_workbook(file_path, data_only=True, read_only=True)
    selected_sheets = wb.sheetnames[sheets_slice]
    data_per_col: Dict[int, List[float]] = {col: [] for col in col_indices}
    for sheet_name in selected_sheets:
        ws = wb[sheet_name]
        for col in col_indices:
            for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=col, max_col=col, values_only=True):
                v = row[0]
                if isinstance(v, (int, float)) and not (isinstance(v, float) and math.isnan(v)):
                    data_per_col[col].append(float(v))
    wb.close()
    return {col: np.array(vals, dtype=float) if len(vals) > 0 else np.array([], dtype=float)
            for col, vals in data_per_col.items()}

def percentile_threshold_for_col(col_data: np.ndarray, top_percent: float) -> Optional[float]:
    if col_data.size == 0:
        return None
    return float(np.percentile(col_data, 100.0 - top_percent))

def describe_array(arr: np.ndarray) -> str:
    if arr.size == 0:
        return "size=0"
    q = np.percentile(arr, [0, 25, 50, 75, 100])
    return f"size={arr.size} min={q[0]:.6f} p25={q[1]:.6f} p50={q[2]:.6f} p75={q[3]:.6f} max={q[4]:.6f}"

# -------------------------------
# Graph score helpers
# -------------------------------

def fetch_graph_scores(driver, gid: str, i: int) -> np.ndarray:
    """
    Fetch all o.sim_score[i] for a gid. Used for graph-based percentile threshold and diagnostics.
    """
    query = """
    MATCH (n {gid: $gid})-[:paragraph]->(m)
    WHERE NOT n:Summary AND NOT m:Summary
    MATCH (m)-[s]-(o)
    WHERE NOT o:Summary AND TYPE(s) <> 'paragraph'
    WITH o.sim_score[$i] AS score
    WHERE score IS NOT NULL
    RETURN score
    """
    with driver.session() as session:
        rows = session.run(query, gid=gid, i=i).data()
    scores = [float(r["score"]) for r in rows if r.get("score") is not None]
    return np.array(scores, dtype=float)

# -------------------------------
# Neo4j link_context with threshold
# -------------------------------

def link_context_neo4j(driver, gid: str, i: int, sim_score_threshold: float, comparison_op: str = COMPARISON_OP) -> List[str]:
    cont: List[str] = []
    if comparison_op not in (">", ">="):
        comparison_op = ">="
    retrieve_query = f"""
        MATCH (n)
        WHERE n.gid = $gid AND NOT n:Summary
        MATCH (n)-[r:paragraph]->(m)
        WHERE NOT m:Summary
        MATCH (m)-[s]-(o)
        WHERE NOT o:Summary AND TYPE(s) <> 'paragraph' AND o.sim_score[$i] {comparison_op} $thr
        RETURN n.id AS NodeId1, 
               m.id AS Mid, 
               TYPE(r) AS paragraphType, 
               collect(DISTINCT {{RelationType: type(s), Oid: o.id}}) AS Connections
    """
    with driver.session() as session:
        res = session.run(
            retrieve_query,
            gid=gid,
            i=i,
            thr=sim_score_threshold
        ).data()
    for r in res:
        connections = r.get("Connections", [])
        for ind, connection in enumerate(connections):
            cont.append(
                "Reference " + str(ind) + ": " + r["NodeId1"] + "has the reference that" + r["Mid"] + connection["RelationType"] + connection["Oid"]
            )
    return cont

# -------------------------------
# Core logic
# -------------------------------

def build_user_two(question: str, last_response: str, references_texts: List[str]) -> str:
    # question / last_response는 옵션. references_texts만으로도 토큰 계산됨.
    return (
            ("the question is: " + (question or "")) +
            ("the last response of it is: " + (last_response or "")) +
            "the references are: " + "".join(references_texts)
    )

def count_total_tokens_for_cont(cont: List[str], question: str, last_response: str, system_prompt: str, model: str) -> int:
    user_content = build_user_two(question, last_response, cont)
    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": user_content},
    ]
    return count_chat_tokens(messages, model=model)

def find_percent_match_for_gid(
        driver,
        gid: str,
        gid_excel_col: int,
        preloaded_cols_data: Dict[int, np.ndarray],
        sim_index: int,
        target_tokens: int,
        question: str,
        last_response: str,
        system_prompt: str,
        model: str,
        percent_start: float,
        percent_end: float,
        percent_step: float,
        token_tolerance: int,
) -> Tuple[float, int]:

    # Base tokens (no references), 측정값 (비교에서 cont 토큰만 분리하기 위해 사용)
    base_tokens_measured = count_total_tokens_for_cont([], question, last_response, system_prompt, model)

    # 목표는 '링크 컨텍스트로 추가되는 토큰'이 target_cont_tokens에 근접하도록
    target_cont_tokens = target_tokens - FIXED_BASE_TOKENS
    if DEBUG:
        print(f"[DEBUG] GID={gid} target_total={target_tokens} FIXED_BASE={FIXED_BASE_TOKENS} target_cont={target_cont_tokens} base_measured={base_tokens_measured}")

    # p=0 (컨텍스트 없음)일 때 cont 토큰은 0
    best_p = 0.0
    best_diff = abs(0 - target_cont_tokens)

    if best_diff <= token_tolerance:
        if DEBUG:
            print(f"[DEBUG] GID={gid} cont=0 within tolerance. Returning p=0.0")
        return (0.0, int(round(best_diff)))

    # Prepare Excel and Graph stats
    col_data = preloaded_cols_data.get(gid_excel_col, np.array([], dtype=float))
    graph_scores = fetch_graph_scores(driver, gid, sim_index)

    if DEBUG:
        print(f"[DEBUG] GID={gid} Excel col stats: {describe_array(col_data)}")
        print(f"[DEBUG] GID={gid} Graph score stats: {describe_array(graph_scores)}")

    p = max(percent_start, 0.0)
    end_p = min(percent_end, 100.0)
    debug_p_logs = 0

    while p <= end_p + 1e-12:
        # threshold from graph or excel
        thr_src = "graph" if (USE_GRAPH_PERCENTILE and graph_scores.size > 0) else "excel"
        if thr_src == "graph":
            thr = float(np.percentile(graph_scores, 100.0 - p))
            thr_adj = thr - THRESHOLD_EPS if COMPARISON_OP == ">" else thr
        else:
            thr = percentile_threshold_for_col(col_data, p)
            if thr is None:
                p += percent_step
                continue
            thr_adj = thr - THRESHOLD_EPS if COMPARISON_OP == ">" else thr

        cont = link_context_neo4j(driver, gid, sim_index, thr_adj, comparison_op=COMPARISON_OP)
        total_tokens = count_total_tokens_for_cont(cont, question, last_response, system_prompt, model)
        cont_tokens = max(0, total_tokens - base_tokens_measured)  # 컨텍스트로 추가된 토큰
        diff = abs(cont_tokens - target_cont_tokens)

        if DEBUG and debug_p_logs < DEBUG_MAX_P_LOGS_PER_SAMPLE:
            print(f"[DEBUG] GID={gid} p={p:.2f} src={thr_src} thr={thr:.6f} thr_adj={thr_adj:.6f} "
                  f"cont_len={len(cont)} total={total_tokens} cont_tokens={cont_tokens} diff={diff}")
            debug_p_logs += 1

        if diff <= token_tolerance:
            if DEBUG:
                print(f"[DEBUG] GID={gid} Matched within tolerance at p={p:.2f}")
            return (round(p, 2), int(round(diff)))

        if diff < best_diff:
            best_diff = diff
            best_p = p

        p += percent_step

    if DEBUG:
        print(f"[DEBUG] GID={gid} No exact match. Returning closest p={best_p:.2f} (best_diff={best_diff})")

    if best_diff > 1000:
        return (100.0, int(round(best_diff)))

    return (round(best_p, 2), int(round(best_diff)))

# -------------------------------
# Main
# -------------------------------

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--out_csv",
        type=str,
        default="19000_simpct_list.csv",
        help="Path to write results CSV (default: 19000_simpct_list.csv).",
    )
    args = parser.parse_args()

    print("[INFO] Loading GIDs from Excel...")
    gids, gid_to_excel_col = load_gids_from_excel(EXCEL_PATH, EXCEL_ROW_INDEX, EXCEL_START_COL_INDEX, EXCEL_GID_COUNT)
    print(f"[INFO] Loaded {len(gids)} GIDs.")

    # 고정 샘플 생성: (gid_idx, FIXED_TARGET_TOKENS) for gid_idx in 0..len(gids)-1
    print("[INFO] Generating fixed samples (no Sobol)...")
    sample_pairs: List[Tuple[int, int]] = [(i, FIXED_TARGET_TOKENS) for i in range(len(gids))]
    print(f"[INFO] Generated {len(sample_pairs)} fixed samples (gid_idx=0..{len(gids)-1}, target_tokens={FIXED_TARGET_TOKENS}).")

    print("[INFO] Preloading Excel column data for percentile thresholds...")
    all_excel_cols = [gid_to_excel_col[g] for g in gids]
    preloaded_cols_data = preload_excel_columns(
        file_path=EXCEL_PATH,
        col_indices=all_excel_cols,
        min_row=OPENPYXL_MIN_ROW,
        max_row=OPENPYXL_MAX_ROW,
        sheets_slice=OPENPYXL_SHEETS_SLICE
    )
    print(f"[INFO] Preloaded columns: {len(preloaded_cols_data)}")

    print(f"[INFO] Connecting to Neo4j at {NEO4J_URI} ...")
    driver = GraphDatabase.driver(NEO4J_URI, auth=(NEO4J_USER, NEO4J_PASSWORD))
    print("[INFO] Neo4j driver created.")

    # results rows: [gid_index (0-based), gid, token_total, simscore_percentage, adj_token, round_adj_token]
    results: List[List[object]] = []
    question_cache: Dict[str, str] = {}

    try:
        for idx, (gid_idx, target_tokens) in enumerate(sample_pairs):
            gid_idx = max(0, min(gid_idx, len(gids) - 1))

            gid = gids[gid_idx]
            gid_excel_col = gid_to_excel_col.get(gid)

            if gid not in question_cache:
                question_cache[gid] = get_question_for_gid(gid)
            question_text = question_cache[gid]

            print(f"\n[INFO] Sample {idx+1}/{len(sample_pairs)}: gid_idx={gid_idx} GID={gid} target_tokens={target_tokens} excel_col={gid_excel_col} use_graph_pct={USE_GRAPH_PERCENTILE} cmp={COMPARISON_OP}")

            match_pct, diff = find_percent_match_for_gid(
                driver=driver,
                gid=gid,
                gid_excel_col=gid_excel_col,
                preloaded_cols_data=preloaded_cols_data,
                sim_index=SIM_SCORE_INDEX,
                target_tokens=target_tokens,
                question=question_text,
                last_response=LAST_RESPONSE,
                system_prompt=SYS_PROMPT_TWO,
                model=MODEL_NAME,
                percent_start=PERCENT_START,
                percent_end=PERCENT_END,
                percent_step=PERCENT_STEP,
                token_tolerance=TOKEN_TOLERANCE,
            )

            # diff는 '컨텍스트 토큰'의 오차
            # 총 토큰 관점에서의 근사값
            adj_token = target_tokens - diff
            print(f"[INFO] Result for gid_idx={gid_idx} {gid}: percent={match_pct}, diff(cont)={diff}, adj_total_tokens≈{adj_token}")

            results.append([gid_idx, gid, target_tokens, match_pct, adj_token, round(adj_token, -3)])
    finally:
        driver.close()
        print("\n[INFO] Neo4j driver closed.")

        # Save results to CSV
        df = pd.DataFrame(results, columns=["gid_index", "gid", "token_total", "simscore_percentage", "adj_token", "round_adj_token"])
        df.to_csv(args.out_csv, index=False, encoding="utf-8-sig")
        print(f"\n[RESULT] Saved {len(df)} rows to CSV: {args.out_csv}")

        # Also print a preview
        print(df.head(10).to_string(index=False))


if __name__ == "__main__":
    main()